"""내보내는 엑셀 한 벌을 만든다.

화면마다 담는 것이 다르지만 꼴은 하나다. 모든 시트가 같은 자리에서 시작하고
같은 줄 높이를 쓴다. 시트를 넘겨도 눈이 흔들리지 않게 하려는 것이다.

  1행  가름 줄
  2행  제목 띠
  3행  안내 띠
  4행  비움(내역 시트만 필터 합계)
  5행  머리글
  6행~ 본문

빛깔은 앱의 기본 벌(쪽빛)을 그대로 쓴다. 밝은 쪽 한 벌로 못 박는다.
받는 쪽의 화면 설정을 따라갈 까닭이 없다.

손으로 짠 xlsx.py와 달리 openpyxl을 쓴다. 머리글 고정과 필터와 그림은
그쪽으로는 만들 수 없다.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.xlsxchart import 가로막대, 세로막대

# ── 쪽빛. frontend/src/utils/palettes.ts의 기본 벌과 같아야 한다. ──
PRIMARY = "5B5FEF"
PRIMARY_D = "4548D6"
PRIMARY_L = "E8E9FD"
TEAL = "00A89F"
CORAL = "E64842"
ZERO = "9E7DB6"
INK = "212529"
SUB = "6C757D"
BORDER = "E9ECEF"
BAND = "FAFAFE"

F = "맑은 고딕"
MONEY = "#,##0;[Red]-#,##0"
# 적은 그대로 부호를 세운다. 들어온 것은 +, 나간 것은 −, 0은 0이다.
# 빛깔은 셀 글꼴로 준다. 엑셀의 서식 빛깔은 여덟 가지뿐이라 앱의 벌을 못 쓴다.
부호돈 = '+#,##0;-#,##0;0'
들어옴글꼴 = Font(name=F, size=10, color=TEAL)
나감글꼴 = Font(name=F, size=10, color=CORAL)
영글꼴 = Font(name=F, size=10, color=ZERO)
DAYFMT = "yyyy-mm-dd"
PCT = "0.0%"
건 = '#,##0"건"'

제목행 = 2
안내행 = 3
머리행 = 5
첫줄 = 6
높이 = {1: 9, 2: 32, 3: 24, 4: 24, 5: 26}
본문높이 = 20
바닥 = 60

제목글꼴 = Font(name=F, size=14, color="FFFFFF", bold=True)
띠글꼴 = Font(name=F, size=11, color="FFFFFF", bold=True)
머리글꼴 = Font(name=F, size=10, color="FFFFFF", bold=True)
본문글꼴 = Font(name=F, size=10, color=INK)
굵은글꼴 = Font(name=F, size=10, color=INK, bold=True)
안내글꼴 = Font(name=F, size=10, color=PRIMARY_D)
라벨글꼴 = Font(name=F, size=10, color=SUB)
숨은글꼴 = Font(name=F, size=10, color="FFFFFF")

제목칠 = PatternFill("solid", fgColor=PRIMARY)
띠칠 = PatternFill("solid", fgColor=TEAL)
머리칠 = PatternFill("solid", fgColor=PRIMARY_D)
안내칠 = PatternFill("solid", fgColor=PRIMARY_L)
줄무늬칠 = PatternFill("solid", fgColor=BAND)
소계칠 = PatternFill("solid", fgColor=PRIMARY_L)

밑선 = Border(bottom=Side(style="thin", color=BORDER))
왼 = Alignment(horizontal="left", vertical="center")
오른 = Alignment(horizontal="right", vertical="center")
가운데 = Alignment(horizontal="center", vertical="center")
들여 = Alignment(horizontal="left", vertical="center", indent=1)

요일글 = "월화수목금토일"


def 원(n: float) -> str:
    return f"{round(n):,}원"


def 조사(말: str, 받침: str, 민낯: str) -> str:
    """받침을 보고 조사를 선택한다. 앱의 narrowOptions와 같은 계산이다."""
    ch = (말 or "").strip()[-1:]
    if not ch:
        return 민낯
    code = ord(ch)
    if 0xAC00 <= code <= 0xD7A3:
        return 받침 if (code - 0xAC00) % 28 else 민낯
    return 받침


def 달글(ym: str) -> str:
    y, m = ym.split("-")[:2]
    return f"{y}년 {int(m)}월"


def 날글(d: date | str) -> str:
    if isinstance(d, str):
        d = datetime.strptime(d[:10], "%Y-%m-%d").date()
    return f"{d.year}. {d.month}. {d.day}."


def 날요일글(d: date | str) -> str:
    """2026. 10. 13. (화). 앱이 묶음 목록에서 쓰는 꼴과 같다."""
    if isinstance(d, str):
        d = datetime.strptime(d[:10], "%Y-%m-%d").date()
    return f"{날글(d)} ({요일글[d.weekday()]})"


class 장:
    """시트 한 장. 자리와 높이를 한 벌로 맞춘다."""

    def __init__(self, wb: Workbook, 이름: str, 폭들: Sequence[float]):
        self.ws = wb.create_sheet(이름)
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100
        ws.sheet_view.topLeftCell = "A1"
        ws.sheet_format.defaultRowHeight = 본문높이
        ws.sheet_format.customHeight = True
        ws.column_dimensions["A"].width = 3
        for 행, h in 높이.items():
            ws.row_dimensions[행].height = h
        for i, w in enumerate(폭들, start=2):
            ws.column_dimensions[get_column_letter(i)].width = w
        self.칸수 = len(폭들)
        self.r = 첫줄

    # ── 띠 ────────────────────────────────────────────────
    def 띠(self, 행: int, 글, 칠: PatternFill, 글꼴: Font, 칸수: int | None = None) -> None:
        """빛깔 띠 한 줄.

        병합하지 않는다. 병합한 칸은 넘치는 글자를 잘라 버리지만, 빈 칸 위로는
        글자가 그대로 흘러간다. 칠만 하면 보기에는 한 덩이다.
        """
        n = 칸수 or self.칸수
        for i in range(2, 2 + n):
            self.ws.cell(row=행, column=i).fill = 칠
        c = self.ws.cell(row=행, column=2, value=글)
        c.font = 글꼴
        c.alignment = 들여

    def 머리(self, 제목, 안내: str) -> None:
        self.띠(제목행, 제목, 제목칠, 제목글꼴)
        self.띠(안내행, 안내, 안내칠, 안내글꼴)

    def 구역(self, 글: str) -> None:
        self.띠(self.r, 글, 띠칠, 띠글꼴)
        self.r += 1

    def 머리줄(self, 이름들: Sequence[str], 행: int | None = None) -> None:
        행 = 머리행 if 행 is None else 행
        for i, h in enumerate(이름들):
            c = self.ws.cell(row=행, column=2 + i, value=h)
            c.font = 머리글꼴
            c.fill = 머리칠
            c.alignment = 가운데

    # ── 줄 ────────────────────────────────────────────────
    def 줄(
        self,
        값들: Sequence[Any],
        꼴: dict[int, str] | None = None,
        맞춤: dict[int, Alignment] | None = None,
        무늬: bool = False,
        글꼴: Font | None = None,
        칠: PatternFill | None = None,
    ) -> int:
        꼴 = 꼴 or {}
        맞춤 = 맞춤 or {}
        r = self.r
        for i, v in enumerate(값들):
            c = self.ws.cell(row=r, column=2 + i, value=v)
            c.font = 글꼴 or 본문글꼴
            c.border = 밑선
            c.alignment = 맞춤.get(i, 왼)
            if i in 꼴:
                c.number_format = 꼴[i]
            if 칠 is not None:
                c.fill = 칠
            elif 무늬:
                c.fill = 줄무늬칠
        self.r += 1
        return r

    def 라벨값(self, 이름: str, 값: Any, 꼴: str | None = None) -> None:
        r = self.r
        a = self.ws.cell(row=r, column=2, value=이름)
        a.font = 라벨글꼴
        a.alignment = 들여
        a.border = 밑선
        b = self.ws.cell(row=r, column=3, value=값)
        b.font = 굵은글꼴 if 꼴 else 본문글꼴
        b.alignment = 오른 if 꼴 else 왼
        b.border = 밑선
        if 꼴:
            b.number_format = 꼴
        for i in range(4, 2 + self.칸수):
            self.ws.cell(row=r, column=i).border = 밑선
        self.r += 1

    def 비었다(self, 말: str) -> None:
        c = self.ws.cell(row=self.r, column=2, value=말)
        c.font = 라벨글꼴
        self.r += 1

    def 거르개(self) -> None:
        """머리글로 거를 수 있게 한다. 줄이 없으면 걸지 않는다."""
        if self.r <= 첫줄:
            return
        끝 = get_column_letter(1 + self.칸수)
        self.ws.auto_filter.ref = f"B{머리행}:{끝}{self.r - 1}"

    def 못박기(self, 가로까지: str = "B") -> None:
        self.ws.freeze_panes = f"{가로까지}{첫줄}"


def 묶음만들기() -> Workbook:
    wb = Workbook()
    wb.remove(wb["Sheet"])
    return wb


def 마무리(wb: Workbook) -> bytes:
    """빈 자리까지 같은 높이로 깔고 첫 장을 연다."""
    for ws in wb:
        for 행 in range(첫줄, max(ws.max_row, 바닥) + 1):
            ws.row_dimensions[행].height = 본문높이
    wb.active = 0
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ══ 자주 쓰는 장 ═══════════════════════════════════════════

def 표제(화면: str) -> CellRichText:
    """돈을 쓰다: "지출 내역"

    서비스 이름은 그대로 세우고 화면 이름만 기울여 한 치수 줄인다.
    둘이 같은 글씨면 어디까지가 이름인지 읽는 데 한 박자가 걸린다.
    """
    return CellRichText(
        TextBlock(InlineFont(rFont=F, sz=14, b=True, color="FFFFFFFF"), "돈을 쓰다: "),
        TextBlock(InlineFont(rFont=F, sz=12, b=True, i=True, color="FFFFFFFF"), f'"{화면}"'),
    )


def 요약장(wb: Workbook, 화면: str, 기간: str, 조건: Sequence[tuple[str, str]],
         요약: Sequence[tuple[str, Any, str]]) -> 장:
    s = 장(wb, "① 요약", [26, 20, 14])
    s.머리(표제(화면), 기간)
    s.r = 첫줄
    s.구역("조회 조건")
    for 이름, 값 in 조건:
        s.라벨값(이름, 값)
    s.r += 1
    s.구역("요약")
    for 이름, 값, 꼴 in 요약:
        s.라벨값(이름, 값, 꼴)
    return s


내역머리 = ["구분", "날짜", "요일", "중분류", "소분류", "세분류", "IN/OUT",
         "금액", "내 몫", "결제 수단", "장소", "묶음", "메모"]
내역폭 = [8, 12, 6, 12, 12, 16, 8, 13, 13, 20, 28, 14, 46]
내역꼴 = {1: DAYFMT, 7: 부호돈, 8: 부호돈}
내역맞춤 = {0: 가운데, 1: 가운데, 2: 가운데, 3: 가운데, 4: 가운데, 6: 가운데, 7: 오른, 8: 오른}


def 내역장(wb: Workbook, 안내: str, 줄들: Iterable[dict], 번호: str = "②",
        덧머리: Sequence[str] = (), 덧폭: Sequence[float] = (),
        덧값=None) -> 장:
    """지출, 대기, 정기를 한 꼴로 담는다. 첫 칸이 그 구분이다.

    어느 화면에서 내보내든 앞의 열세 칸은 같다. 화면마다 더 담을 것이 있으면
    뒤에 덧붙인다. 앞을 흔들면 두 파일을 나란히 놓고 볼 수가 없다.
    """
    # 어느 화면에서 내보내든 이름도 제목도 "내역" 하나다. 어느 화면 것인지는
    # 파일 이름과 ① 요약의 표제가 말한다.
    s = 장(wb, f"{번호} 내역", list(내역폭) + list(덧폭))
    s.머리("내역", 안내)
    줄들 = list(줄들)
    끝행 = 첫줄 - 1 + len(줄들)

    c = s.ws.cell(row=4, column=8, value="필터 합계")
    c.font = Font(name=F, size=10, color=SUB, bold=True)
    c.alignment = 오른
    for col in (9, 10):
        열 = get_column_letter(col)
        c = s.ws.cell(row=4, column=col, value=f"=SUBTOTAL(109,{열}{첫줄}:{열}{max(끝행, 첫줄)})")
        c.font = Font(name=F, size=11, color=PRIMARY_D, bold=True)
        c.number_format = MONEY
        c.alignment = 오른
    c = s.ws.cell(row=4, column=11, value=f'=SUBTOTAL(103,B{첫줄}:B{max(끝행, 첫줄)})&"건"')
    c.font = Font(name=F, size=10, color=SUB)
    c.alignment = 왼

    s.머리줄(list(내역머리) + list(덧머리))
    for 칸 in (9, 10):
        c = s.ws.cell(row=4, column=칸)
        c.number_format = 부호돈
    for j, x in enumerate(줄들):
        # 멈춘 정기 규칙은 다음에 생길 날이 없다. 날짜 칸을 비운다.
        d = x["tx_date"]
        if isinstance(d, str):
            d = datetime.strptime(d[:10], "%Y-%m-%d").date()
        elif isinstance(d, datetime):
            d = d.date()
        # 적힌 그대로 둔다. 들어온 것은 +, 나간 것은 −다.
        부호 = -1 if x["inout"] == -1 else 1
        몫 = float(x["amount"]) - float(x.get("split_sum") or 0)
        r = s.줄(
            [x.get("kind") or "지출", d, 요일글[d.weekday()] if d else None,
             x.get("cat1_name"), x.get("cat2_name"), x.get("cat3_name"),
             "IN" if x["inout"] == 1 else "OUT",
             float(x["amount"]) * 부호, 몫 * 부호,
             x.get("method_name"), x.get("place_name"), x.get("group_name"), x.get("memo")]
            + (list(덧값(x)) if 덧값 else []),
            꼴=내역꼴, 맞춤=내역맞춤, 무늬=(j % 2 == 1),
        )
        # 빛깔은 앱과 같다. 들어옴은 청록, 나감은 산호, 0은 제3의 빛깔이다.
        들어옴 = x["inout"] == 1
        글꼴 = 들어옴글꼴 if 들어옴 else 나감글꼴
        # 금액(9칸)과 내 몫(10칸). 0이면 들어온 것도 나간 것도 아니다.
        for 칸 in (9, 10):
            값 = s.ws.cell(row=r, column=칸).value
            s.ws.cell(row=r, column=칸).font = (
                영글꼴 if isinstance(값, (int, float)) and 값 == 0 else 글꼴)
        # IN/OUT 딱지(8칸)도 같은 빛깔로 세운다.
        s.ws.cell(row=r, column=8).font = Font(
            name=F, size=10, color=TEAL if 들어옴 else CORAL, bold=True)
    if not 줄들:
        s.비었다("조회 조건에 걸린 내역이 없습니다.")
    s.거르개()
    s.못박기("E")
    return s


def 분류장(wb: Workbook, 묶: dict[str, list[dict]], 지출: float, 번호: str = "③") -> 장:
    s = 장(wb, f"{번호} 분류별 집계", [16, 18, 14, 12, 10, 14])
    s.머리("분류별 집계", "진한 줄이 중분류 합계입니다. 비중은 지출 합계 대비입니다.")
    s.머리줄(["중분류", "소분류", "합계", "비중", "건수", "건당 평균"])
    꼴 = {2: MONEY, 3: PCT, 4: "#,##0", 5: MONEY}
    맞춤 = {2: 오른, 3: 오른, 4: 오른, 5: 오른}
    처음 = s.r
    for c1, 목록 in sorted(묶.items(), key=lambda kv: -sum(float(t["amt"]) for t in kv[1])):
        소계 = sum(float(t["amt"]) for t in 목록)
        건수 = sum(int(t["n"]) for t in 목록)
        s.줄([c1, None, 소계, 소계 / 지출 if 지출 else 0, 건수, round(소계 / 건수) if 건수 else 0],
             꼴=꼴, 맞춤=맞춤, 글꼴=굵은글꼴, 칠=소계칠)
        for t in 목록:
            amt = float(t["amt"])
            n = int(t["n"])
            r = s.줄([c1, t["cat2_name"], amt, amt / 지출 if 지출 else 0, n, round(amt / n) if n else 0],
                    꼴=꼴, 맞춤={1: 들여, **맞춤})
            # 중분류를 아래 줄에도 담되 바탕과 같은 빛깔로 적는다. 눈에는 안
            # 보이지만 머리글로 거를 때 딸린 줄이 함께 남는다.
            s.ws.cell(row=r, column=2).font = 숨은글꼴
    if 처음 < s.r:
        s.ws.conditional_formatting.add(
            f"E{처음}:E{s.r - 1}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=0.45,
                        color=PRIMARY, showValue=True, minLength=None, maxLength=None))
    else:
        s.비었다("집계할 지출이 없습니다.")
    s.거르개()
    s.못박기()
    return s


def 그림장(wb: Workbook, 그림들: Sequence[tuple[str, bytes]], 번호: str = "④") -> 장:
    """그린 그림을 얹는다. 4배로 그려 1/4로 줄여 넣어 글자가 또렷하다."""
    s = 장(wb, f"{번호} 그림", [14] * 8)
    s.머리("그림", "앱의 씀씀이와 같은 결로 그립니다. 값은 조회 조건을 그대로 따릅니다.")
    행 = 첫줄
    for _, png in 그림들:
        im = XLImage(BytesIO(png))
        im.width, im.height = im.width / 4, im.height / 4
        s.ws.add_image(im, f"B{행}")
        행 += int(im.height / 26.67) + 2
    s.r = 행
    if not 그림들:
        s.비었다("그릴 것이 없습니다.")
    return s


def 영수증장(wb: Workbook, 쪽들: Sequence[bytes], 번호: str = "⑦") -> 장 | None:
    """앱이 그린 영수증을 그대로 얹는다.

    길면 잘라서 옆으로 나란히 세운다. 한 장으로 길게 두면 시트가 수천 줄이
    되어 넘길 때마다 자리가 흔들린다. 쪽마다 같은 줄에서 시작하므로 가로로만
    늘어나고 세로는 늘 같다.
    """
    if not 쪽들:
        return None
    # 칸을 잘게 썰어 쪽 사이가 벌어지지 않게 한다. 칸 하나가 넓으면 그림
    # 하나를 세울 때마다 쓰지도 않는 자리가 따라 넓어진다.
    칸폭 = 2.4
    칸px = 칸폭 * 7
    쪽폭 = max((XLImage(BytesIO(p)).width / 2) for p in 쪽들)
    쪽칸 = int(-(-쪽폭 // 칸px)) + 2  # 그림이 차지하는 칸 + 사이 두 칸
    s = 장(wb, f"{번호} 영수증", [칸폭] * (쪽칸 * len(쪽들) + 4))
    s.머리("영수증", "앱에서 보는 영수증과 같은 그림입니다. 길어서 잘라 옆으로 세웠습니다.")
    for i, png in enumerate(쪽들):
        im = XLImage(BytesIO(png))
        im.width, im.height = im.width / 2, im.height / 2
        s.ws.add_image(im, f"{get_column_letter(2 + i * 쪽칸)}{첫줄}")
    s.r = 첫줄
    return s


def 도전장(wb: Workbook, 줄들: Sequence[dict], 번호: str = "⑤") -> 장:
    s = 장(wb, f"{번호} 안쓴이 도전", [12, 16, 16, 14, 14, 14, 10, 34])
    s.머리("안쓴이 도전",
          "조회 기간에 걸친 달의 목표와 실제 지출을 견줍니다. 남은 금액이 음수면 넘긴 것입니다.")
    s.머리줄(["연월", "중분류", "소분류", "목표", "지출", "남은 금액", "달성률", "메모"])
    for j, g in enumerate(줄들):
        목표, 실제 = float(g["목표"]), float(g["실제"])
        r = s.줄([달글(g["ym"]), g["cat1_name"], g["cat2_name"], 목표, 실제, 목표 - 실제,
                실제 / 목표 if 목표 else 0, g["memo"]],
               꼴={3: MONEY, 4: MONEY, 5: MONEY, 6: PCT},
               맞춤={0: 가운데, 3: 오른, 4: 오른, 5: 오른, 6: 오른},
               무늬=(j % 2 == 1))
        빛 = CORAL if 실제 > 목표 else TEAL
        for i in (7, 8):
            s.ws.cell(row=r, column=i).font = Font(name=F, size=10, color=빛, bold=True)
    if not 줄들:
        s.비었다("조회 기간에 걸린 도전이 없습니다.")
    s.거르개()
    s.못박기()
    return s


def 잔소리장(wb: Workbook, 말들: Sequence[tuple[str, str, str]], 번호: str = "⑥") -> 장:
    s = 장(wb, f"{번호} 잔소리", [10, 62, 42])
    s.머리("잔소리", "앱의 잔소리와 같은 기준으로 적습니다. :-( 는 살펴볼 것, ;-) 는 잘한 것입니다.")
    s.머리줄(["구분", "한 줄", "근거"])
    for j, (수준, 한줄, 근거) in enumerate(말들):
        r = s.줄([수준, 한줄, 근거], 맞춤={0: 가운데}, 무늬=(j % 2 == 1))
        s.ws.cell(row=r, column=2).font = Font(
            name=F, size=10, bold=True, color=CORAL if 수준 == ":-(" else TEAL)
    if not 말들:
        s.비었다("잔소리할 게 없습니다.")
    s.거르개()
    s.못박기()
    return s


__all__ = [
    "묶음만들기", "마무리", "장", "요약장", "내역장", "분류장", "그림장", "도전장", "잔소리장",
    "영수증장",
    "MONEY", "DAYFMT", "PCT", "건", "왼", "오른", "가운데", "들여",
    "원", "조사", "달글", "날글", "날요일글", "세로막대", "가로막대",
]
